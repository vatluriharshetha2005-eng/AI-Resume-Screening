from io import BytesIO
from pathlib import Path


def extract_text(file_bytes, filename):
    extension = Path(filename).suffix.lower()

    if extension == ".txt":
        return file_bytes.decode("utf-8", errors="ignore")

    elif extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(file_bytes))
        text = ""

        for page in reader.pages:
            text += page.extract_text() or ""

        return text

    elif extension == ".docx":
        from docx import Document

        document = Document(BytesIO(file_bytes))
        text = "\n".join(
            paragraph.text
            for paragraph in document.paragraphs
        )

        for table in document.tables:
            for row in table.rows:
                text += "\n" + " ".join(
                    cell.text for cell in row.cells
                )

        return text

    else:
        raise ValueError("Unsupported file format")